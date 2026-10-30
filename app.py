import streamlit as st
import pandas as pd
import os
import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from io import BytesIO
from streamlit_gsheets import GSheetsConnection

ADMIN_PASSWORD = "admin123"

# Упрощенная функция генерации объективки без локальных фото
def generate_docx_anketa(row):
    doc = Document()
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.6)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("МАЪЛУМОТНОМА")
    run_title.bold = True
    run_title.font.size = Pt(13)
    
    p_fio = doc.add_paragraph()
    p_fio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fio.paragraph_format.space_after = Pt(20)
    run_fio = p_fio.add_run(str(row.get("ФИО", "-")))
    run_fio.bold = True
    run_fio.underline = True
    run_fio.font.size = Pt(13)
    
    def clean_val(val):
        if pd.isna(val) or str(val).strip().lower() == "nan" or not str(val).strip():
            return "-"
        return str(val).strip()

    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    fields = [
        ("Туғилган йили:", row.get("Дата рождения")),
        ("Туғилган жойи:", row.get("Место рождения")),
        ("Миллати:", row.get("Национальность")),
        ("Партиявийлиги:", "йўқ"),
        ("Маълумоти:", f"{clean_val(row.get('Уровень образования'))} ({clean_val(row.get('Специальность'))})"),
        ("Тамомлаган:", row.get("Учебное заведение")),
        ("Маълумоти бўйича мутахассислиги:", row.get('Специальность')),
        ("Илмий даражаси / унвони:", "йўқ"),
        ("Қайси чет тилларини билади:", f"русский язык и английский язык ({clean_val(row.get('Уровень английского'))})"),
        ("Давлат мукофотлари:", "йўқ"),
        ("Жорий лавозими:", row.get("Должность"))
    ]
    
    for label, val in fields:
        r = p.add_run(f"{label:<32}")
        r.bold = True
        p.add_run(f"{clean_val(val)}\n")
    
    p_work_head = doc.add_paragraph()
    p_work_head.paragraph_format.space_before = Pt(20)
    p_work_head.paragraph_format.space_after = Pt(8)
    p_work_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_wh = p_work_head.add_run("МЕҲНАТ ФАОЛИЯТИ")
    r_wh.bold = True
    
    w_table = doc.add_table(rows=2, cols=2)
    w_table.style = 'Table Grid'
    w_table.columns[0].width = Inches(1.5)
    w_table.columns[1].width = Inches(5.3)
    
    w_cells_hdr = w_table.rows[0].cells
    w_cells_hdr[0].text = "Йиллар"
    w_cells_hdr[1].text = "Ишлаган жойи ва лавозими"
    
    exp_years = row.get("Стаж работы (лет)", 0)
    exp_str = "0 лет" if pd.isna(exp_years) or str(exp_years).lower() == "nan" else f"{int(exp_years)} лет"
    
    w_cells_body = w_table.rows[1].cells
    w_cells_body[0].text = "Сбор данных"
    w_cells_body[1].text = f"Общий стаж работы составляет {exp_str}. Текущая должность: {clean_val(row.get('Должность'))}"
    
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

st.set_page_config(page_title="Кадровая анкета / Anketa", page_icon="👤", layout="centered")

# Подключение к Google Таблицам через облачный коннектор
try:
    conn = st.connection("gsheets", type=GSheetsConnection)
    df_existing = conn.read(ttl=0)
except Exception as e:
    st.error("Пожалуйста, настройте ссылку на Google Sheets в панели Secrets.")
    df_existing = pd.DataFrame()

# Языковой пакет (Сокращен для примера)
LANG_PACK = {
    "Русский": {"title": "👤 Расширенная анкетa сотрудника", "subtitle": "Заполните форму.", "fio": "ФИО сотрудника *", "submit": "Отправить анкету", "success": "🎉 Данные успешно сохранены!"},
    "O'zbekcha": {"title": "👤 Xodimning anketasi", "subtitle": "Formani to'ldiring.", "fio": "Xodimning F.I.Sh. *", "submit": "Anketani yuborish", "success": "🎉 Saqlandi!"},
    "English": {"title": "👤 Employee Questionnaire", "subtitle": "Fill out the form.", "fio": "Full Name *", "submit": "Submit", "success": "🎉 Saved!"}
}

selected_lang = st.selectbox("🌐 Язык / Til / Language:", ["Русский", "O'zbekcha", "English"])
T = LANG_PACK[selected_lang]

st.title(T["title"])

with st.form(key="employee_form", clear_on_submit=True):
    fio = st.text_input(T["fio"])
    birthday = st.date_input("Дата рождения *", value=datetime.date(1995, 1, 1))
    bplace = st.text_input("Место рождения *")
    nationality = st.text_input("Национальность *")
    phone = st.text_input("Номер телефона *")
    address = st.text_input("Адрес проживания *")
    position = st.text_input("Должность *")
    education = st.selectbox("Образование", ["Среднее", "Высшее"])
    edu_inst = st.text_input("Учебное заведение *")
    specialty = st.text_input("Специальность *")
    english_level = st.selectbox("Английский", ["A1", "B1", "C1"])
    experience = st.number_input("Стаж (лет)", min_value=0, max_value=60, step=1)
    
    submit_button = st.form_submit_button(label=T["submit"])

if submit_button:
    if not fio.strip() or not phone.strip() or not address.strip() or not position.strip():
        st.error("Заполните поля!")
    else:
        new_row = pd.DataFrame([{
            "ФИО": fio.strip(),
            "Дата рождения": birthday.strftime("%d.%m.%Y"),
            "Место рождения": bplace.strip(),
            "Национальность": nationality.strip(),
            "Номер телефона": phone.strip(),
            "Адрес проживания": address.strip(),
            "Должность": position.strip(),
            "Уровень образования": education,
            "Учебное заведение": edu_inst.strip(),
            "Специальность": specialty.strip(),
            "Уровень английского": english_level,
            "Стаж работы (лет)": int(experience)
        }])
        
        updated_df = pd.concat([df_existing, new_row], ignore_index=True)
        conn.update(data=updated_df)
        st.success(T["success"])

st.write("---")
with st.expander("🔐 Панель администратора"):
    input_password = st.text_input("Пароль:", type="password")
    if input_password == ADMIN_PASSWORD:
        st.success("Доступ разрешен")
        if not df_existing.empty:
            st.dataframe(df_existing, use_container_width=True)
            selected_emp = st.selectbox("Сотрудник для выгрузки:", df_existing["ФИО"].tolist())
            if selected_emp:
                emp_data = df_existing[df_existing["ФИО"] == selected_emp].iloc[0].to_dict()
                docx_buffer = generate_docx_anketa(emp_data)
                st.download_button(label=f"📥 Скачать объективку {selected_emp}", data=docx_buffer, file_name=f"Malumotnoma_{selected_emp}.docx")
