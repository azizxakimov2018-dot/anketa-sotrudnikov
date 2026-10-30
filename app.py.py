import streamlit as st
import pandas as pd
import os
import datetime
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from io import BytesIO

# Настройки папок, лимитов и безопасности
FILE_NAME = "employee_data.xlsx"
UPLOAD_FOLDER_PHOTOS = "uploads/photos"
UPLOAD_FOLDER_RESUMES = "uploads/resumes"
MAX_FILE_SIZE_MB = 5
MAX_FILE_SIZE_BYTES = MAX_FILE_SIZE_MB * 1024 * 1024

ADMIN_PASSWORD = "admin123"

os.makedirs(UPLOAD_FOLDER_PHOTOS, exist_ok=True)
os.makedirs(UPLOAD_FOLDER_RESUMES, exist_ok=True)

# Словарь переводов интерфейса (RU, UZ, EN)
LANG_PACK = {
    "Русский": {
        "title": "👤 Расширенная анкетa сотрудника",
        "subtitle": "Заполните форму и загрузите необходимые документы.",
        "sec_personal": "📍 Личные данные, контакты и должность",
        "fio": "ФИО сотрудника *",
        "fio_p": "Иванов Иван Иванович",
        "bday": "Дата рождения *",
        "bplace": "Место рождения *",
        "bplace_p": "г. Самарканд",
        "nationality": "Национальность *",
        "nationality_p": "Узбек / Русский",
        "phone": "Номер телефона *",
        "phone_p": "+7 (999) 123-45-67",
        "address": "Адрес проживания *",
        "address_p": "г. Ташкент, Юнусабадский р-н...",
        "position": "Желаемая / текущая должность *",
        "position_p": "Менеджер по продажам / Инженер",
        "sec_edu": "🎓 Образование, опыт и языки",
        "edu_level": "Уровень образования",
        "edu_vals": ["Среднее", "Среднее специальное", "Высшее (Бакалавриат)", "Высшее (Магистратура)", "Ученая степень"],
        "edu_inst": "Учебное заведение *",
        "edu_inst_p": "ТГТУ / МГУ",
        "specialty": "Специальность *",
        "specialty_p": "Информационные технологии",
        "eng_level": "Уровень английского языка",
        "eng_vals": ["Начальный (A1)", "Элементарный (A2)", "Средний (B1)", "Выше среднего (B2)", "Продвинутый (C1)", "В совершенстве (C2)"],
        "exp": "Общий стаж работы (полных лет)",
        "sec_docs": "📁 Загрузка документов",
        "photo": "Загрузите фото сотрудника (Не более 5 МБ)",
        "resume": "Загрузите резюме",
        "submit": "Отправить анкету в базу",
        "err_fields": "⚠️ Пожалуйста, заполните все текстовые поля, отмеченные звездочкой (*)",
        "err_size": f"❌ Ошибка! Размер фотографии слишком большой. Максимальный размер — 5 МБ.",
        "success": "🎉 Данные успешно сохранены!",
        "admin_title": "🔐 Панель администратора (Кадры / HR)",
        "admin_pass_lbl": "Введите пароль для доступа к базе данных:",
        "admin_ok": "🔓 Доступ разрешен",
        "admin_err": "❌ Неверный пароль!",
        "admin_total": "Всего анкет в базе:",
        "admin_search": "🔍 Быстрый поиск сотрудника по ФИО:",
        "admin_empty": "📂 База данных пока пуста.",
        "admin_download": "📥 Скачать полный файл Excel"
    },
    "O'zbekcha": {
        "title": "👤 Xodimning kengaytirilgan anketasi",
        "subtitle": "Formani to'ldiring va kerakli hujjatlarni yuklang.",
        "sec_personal": "📍 Shaxsiy ma'lumotlar, kontaktlar va lavozim",
        "fio": "Xodimning F.I.Sh. *",
        "fio_p": "Ivanov Ivan Ivanovich",
        "bday": "Tug'ilgan sana *",
        "bplace": "Tug'ilgan joyi *",
        "bplace_p": "Samarqand shahri",
        "nationality": "Millati *",
        "nationality_p": "O'zbek",
        "phone": "Telefon raqami *",
        "phone_p": "+998 (90) 123-45-67",
        "address": "Yashash manzili *",
        "address_p": "Toshkent sh., Yunusobod tumani...",
        "position": "Kutilayotgan / joriy lavozim *",
        "position_p": "Sotuvlar menejeri / Muhandis",
        "sec_edu": "🎓 Ma'lumot, tajriba va tillar",
        "edu_level": "Ma'lumoti darajasi",
        "edu_vals": ["O'rta", "O'rta maxsus", "Oliy (Bakalavriat)", "Oliy (Magistratura)", "Ilmiy daraja"],
        "edu_inst": "O'quv yurti *",
        "edu_inst_p": "TDTU / O'zMU",
        "specialty": "Mutaxassislik *",
        "specialty_p": "Axborot texnologiyalari",
        "eng_level": "Ingliz tili darajasi",
        "eng_vals": ["Boshlang'ich (A1)", "Elementar (A2)", "O'rta (B1)", "O'rtadan yuqori (B2)", "Ilg'or (C1)", "Mukammal (C2)"],
        "exp": "Umumiy ish staji (to'liq yillar)",
        "sec_docs": "📁 Hujjatlarni yuklash",
        "photo": "Xodim fotosuratini yuklang (5 MB dan ko'p bo'lmagan)",
        "resume": "Rezyumeni yuklang",
        "submit": "Anketani bazaga yuborish",
        "err_fields": "⚠️ Iltimos, yulduzcha (*) bilan belgilangan barcha matn maydonlarini to'ldiring",
        "err_size": f"❌ Xato! Fotosurat hajmi juda katta. Maksimal hajm — 5 MB.",
        "success": "🎉 Ma'lumotlar muvaffaqiyatli saqlandi!",
        "admin_title": "🔐 Administrator paneli (Kadrlar / HR)",
        "admin_pass_lbl": "Ma'lumotlar bazasiga kirish uchun parolni kiriting:",
        "admin_ok": "🔓 Ruxsat berildi",
        "admin_err": "❌ Noto'g'ri parol!",
        "admin_total": "Bazadagi jami anketalar:",
        "admin_search": "🔍 F.I.Sh. bo'yicha tezkor qidiruv:",
        "admin_empty": "📂 Ma'lumotlar bazasi hozircha bo'sh.",
        "admin_download": "📥 To'liq Excel faylini yuklab olish"
    },
    "English": {
        "title": "👤 Extended Employee Questionnaire",
        "subtitle": "Please fill out the form and upload the required documents.",
        "sec_personal": "📍 Personal Data, Contacts & Position",
        "fio": "Employee Full Name *",
        "fio_p": "John Doe",
        "bday": "Date of Birth *",
        "phone": "Phone Number *",
        "phone_p": "+1 (555) 019-2834",
        "address": "Residential Address *",
        "address_p": "123 Main St, New York...",
        "position": "Desired / Current Position *",
        "position_p": "Sales Manager / Engineer",
        "sec_edu": "🎓 Education, Experience & Languages",
        "edu_level": "Education Level",
        "edu_vals": ["High School", "Vocational", "Bachelor's Degree", "Master's Degree", "Ph.D. / Academic Degree"],
        "edu_inst": "Educational Institution *",
        "edu_inst_p": "MIT / Stanford",
        "specialty": "Major / Specialty *",
        "specialty_p": "Information Technology",
        "eng_level": "English Language Level",
        "eng_vals": ["Beginner (A1)", "Elementary (A2)", "Intermediate (B1)", "Upper-Intermediate (B2)", "Advanced (C1)", "Proficiency (C2)"],
        "exp": "Total Work Experience (Full Years)",
        "sec_docs": "📁 Document Upload",
        "photo": "Upload Employee Photo (Max 5 MB)",
        "resume": "Upload Resume / CV",
        "submit": "Submit Questionnaire",
        "err_fields": "⚠️ Please fill in all fields marked with an asterisk (*)",
        "err_size": f"❌ Error! Photo file size is too large. Maximum size is 5 MB.",
        "success": "🎉 Data successfully saved!",
        "admin_title": "🔐 Admin Panel (HR / Management)",
        "admin_pass_lbl": "Enter password to access database:",
        "admin_ok": "🔓 Access Granted",
        "admin_err": "❌ Incorrect password!",
        "admin_total": "Total entries in database:",
        "admin_search": "🔍 Quick search by Full Name:",
        "admin_empty": "📂 Database is currently empty.",
        "admin_download": "📥 Download Full Excel File"
    }
}

def save_uploaded_file(uploaded_file, folder, employee_name):
    if uploaded_file is not None:
        ext = os.path.splitext(uploaded_file.name)
        clean_name = "".join(c for c in employee_name if c.isalnum() or c in (' ', '_', '-')).rstrip()
        filename = f"{clean_name}{ext}"
        full_path = os.path.join(folder, filename)
        with open(full_path, "wb") as f:
            f.write(uploaded_file.getbuffer())
        return full_path
    return "Не загружено"

# Функция генерации объективки (МАЪЛУМОТНОМА)
def generate_docx_anketa(row):
    doc = Document()
    
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.6)
        section.bottom_margin = Inches(0.6)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.6)

    style = doc.styles['Normal']
    font = style.font
    font.name = 'Times New Roman'
    font.size = Pt(11)
    
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = p_title.add_run("МАЪЛУМОТНОМА")
    run_title.bold = True
    run_title.font.size = Pt(13)
    
    p_fio = doc.add_paragraph()
    p_fio.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_fio.paragraph_format.space_after = Pt(20)
    run_fio = p_fio.add_run(str(row.get("ФИО", "-")))
    run_fio.bold = True
    run_fio.underline = True
    run_fio.font.size = Pt(13)
    
    main_table = doc.add_table(rows=1, cols=2)
    main_table.autofit = False
    
    for row_item in main_table.rows:
        for cell in row_item.cells:
            tcPr = cell._tc.get_or_add_tcPr()
            tcBorders = parse_xml(f'<w:tcBorders {nsdecls("w")}><w:top w:val="none"/><w:left w:val="none"/><w:bottom w:val="none"/><w:right w:val="none"/></w:tcBorders>')
            tcPr.append(tcBorders)

    main_table.columns[0].width = Inches(5.0)
    main_table.columns[1].width = Inches(1.8)
    
    cell_left = main_table.cell(0, 0)
    cell_right = main_table.cell(0, 1)
    
    photo_path = row.get("Путь к фото", "Не загружено")
    if photo_path != "Не загружено" and os.path.exists(str(photo_path)):
        try:
            p_photo = cell_right.paragraphs[0]
            p_photo.alignment = WD_ALIGN_PARAGRAPH.RIGHT
            p_photo.add_run().add_picture(str(photo_path), width=Inches(1.5))
        except:
            pass
            
    def clean_val(val):
        if pd.isna(val) or str(val).strip().lower() == "nan" or not str(val).strip():
            return "-"
        return str(val).strip()

    def add_info_block(cell, label1, val1, label2, val2):
        p = cell.paragraphs[0] if len(cell.paragraphs) == 1 and cell.paragraphs[0].text == "" else cell.add_paragraph()
        p.paragraph_format.space_after = Pt(4)
        p.paragraph_format.line_spacing = 1.15
        
        r_l1 = p.add_run(f"{label1:<28}")
        r_l1.bold = True
        p.add_run(clean_val(val1) + "\n")
        
        r_l2 = p.add_run(f"{label2:<28}")
        r_l2.bold = True
        p.add_run(clean_val(val2))

    bday_val = row.get("Дата рождения")
    bplace_val = row.get("Место рождения")
    nat_val = row.get("Национальность")
    pos_val = row.get("Должность")
    edu_val = f"{clean_val(row.get('Уровень образования'))} ({clean_val(row.get('Специальность'))})"
    inst_val = row.get("Учебное заведение")
    lang_val = f"русский язык и английский язык ({clean_val(row.get('Уровень английского'))})"
    
    add_info_block(cell_left, "Туғилган йили:", bday_val, "Туғилган жойи:", bplace_val)
    add_info_block(cell_left, "Миллати:", nat_val, "Партиявийлиги:", "йўқ")
    add_info_block(cell_left, "Маълумоти:", edu_val, "Тамомлаган:", inst_val)
    add_info_block(cell_left, "Маълумоти бўйича мутахассислиги:", row.get('Специальность'), "Илмий даражаси / унвони:", "йўқ")
    add_info_block(cell_left, "Қайси чет тилларини билади:", lang_val, "Давлат мукофотлари:", "йўқ")
    
    p_work_head = doc.add_paragraph()
    p_work_head.paragraph_format.space_before = Pt(20)
    p_work_head.paragraph_format.space_after = Pt(8)
    p_work_head.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r_wh = p_work_head.add_run("МЕҲНАТ ФАОЛИЯТИ")
    r_wh.bold = True
    
    w_table = doc.add_table(rows=2, cols=2)
    w_table.style = 'Table Grid'
    w_table.columns[0].width = Inches(1.5)
    w_table.columns[1].width = Inches(5.3)
    
    w_cells_hdr = w_table.rows[0].cells
    w_cells_hdr[0].text = "Йиллар"
    w_cells_hdr[1].text = "Ишлаган жойи ва лавозими"
    w_cells_hdr[0].paragraphs[0].runs[0].font.bold = True
    w_cells_hdr[1].paragraphs[0].runs[0].font.bold = True
    
    exp_years = row.get("Стаж работы (лет)", 0)
    exp_str = "0 лет" if pd.isna(exp_years) or str(exp_years).lower() == "nan" else f"{int(exp_years)} лет"
    
    w_cells_body = w_table.rows[1].cells
    w_cells_body[0].text = "Сбор данных"
    w_cells_body[1].text = f"Общий стаж работы составляет {exp_str}. Текущая должность: {clean_val(pos_val)}"
    
    bio = BytesIO()
    doc.save(bio)
    bio.seek(0)
    return bio

st.set_page_config(page_title="Кадровая анкета / Anketa", page_icon="👤", layout="centered")

selected_lang = st.selectbox("🌐 Выберите язык / Tilni tanlang / Select Language:", ["Русский", "O'zbekcha", "English"])
T = LANG_PACK[selected_lang]

st.title(T["title"])
st.write(T["subtitle"])

# ----------------------------------------------------
# СЕКЦИЯ 1: ФОРМА ДЛЯ СОТРУДНИКОВ
# ----------------------------------------------------
with st.form(key="employee_form", clear_on_submit=True):
    st.subheader(T["sec_personal"])
    fio = st.text_input(T["fio"], placeholder=T["fio_p"])
    
    today = datetime.date.today()
    birthday = st.date_input(
        T["bday"],
        value=datetime.date(1995, 1, 1),
        min_value=datetime.date(today.year - 100, 1, 1),
        max_value=today,
        format="DD.MM.YYYY"
    )
    
    bplace = st.text_input(T["bplace"], placeholder=T["bplace_p"])
    nationality = st.text_input(T["nationality"], placeholder=T["nationality_p"])
    phone = st.text_input(T["phone"], placeholder=T["phone_p"])
    address = st.text_input(T["address"], placeholder=T["address_p"])
    position = st.text_input(T["position"], placeholder=T["position_p"])
    
    st.subheader(T["sec_edu"])
    education = st.selectbox(T["edu_level"], T["edu_vals"])
    edu_inst = st.text_input(T["edu_inst"], placeholder=T["edu_inst_p"])
    specialty = st.text_input(T["specialty"], placeholder=T["specialty_p"])
    english_level = st.selectbox(T["eng_level"], T["eng_vals"])
    experience = st.number_input(T["exp"], min_value=0, max_value=60, value=0, step=1)
    
    st.subheader(T["sec_docs"])
    photo_file = st.file_uploader(T["photo"], type=["jpg", "jpeg", "png"])
    resume_file = st.file_uploader(T["resume"], type=["pdf", "doc", "docx"])
    
    submit_button = st.form_submit_button(label=T["submit"])

if submit_button:
    if not fio.strip() or not phone.strip() or not address.strip() or not position.strip() or not edu_inst.strip() or not specialty.strip() or not bplace.strip() or not nationality.strip():
        st.error(T["err_fields"])
    elif photo_file is not None and photo_file.size > MAX_FILE_SIZE_BYTES:
        st.error(T["err_size"])
    else:
        photo_path = save_uploaded_file(photo_file, UPLOAD_FOLDER_PHOTOS, fio)
        resume_path = save_uploaded_file(resume_file, UPLOAD_FOLDER_RESUMES, fio)
        
        new_data = {
            "ФИО": [fio.strip()],
            "Дата рождения": [birthday.strftime("%d.%m.%Y")],
            "Место рождения": [bplace.strip()],
            "Национальность": [nationality.strip()],
            "Номер телефона": [phone.strip()],
            "Адрес проживания": [address.strip()],
            "Должность": [position.strip()],
            "Уровень образования": [education],
            "Учебное заведение": [edu_inst.strip()],
            "Специальность": [specialty.strip()],
            "Уровень английского": [english_level],
            "Стаж работы (лет)": [int(experience)],
            "Путь к фото": [photo_path],
            "Путь к резюме": [resume_path]
        }
        new_df = pd.DataFrame(new_data)
        
        try:
            if os.path.exists(FILE_NAME):
                old_df = pd.read_excel(FILE_NAME)
                final_df = pd.concat([old_df, new_df], ignore_index=True)
            else:
                final_df = new_df
            
            final_df.to_excel(FILE_NAME, index=False)
            st.success(T["success"])
        except Exception as e:
            st.error(f"❌ Error: {e}")

st.write("---")

# ----------------------------------------------------
# СЕКЦИЯ 2: ПАНЕЛЬ АДМИНИСТРАТОРА
# ----------------------------------------------------
with st.expander(T["admin_title"]):
    input_password = st.text_input(T["admin_pass_lbl"], type="password")
    
    if input_password == ADMIN_PASSWORD:
        st.success(T["admin_ok"])
        
        if os.path.exists(FILE_NAME):
            df_admin = pd.read_excel(FILE_NAME)
            st.metric(T["admin_total"], len(df_admin))
            
            search_query = st.text_input(T["admin_search"])
            if search_query:
                df_filtered = df_admin[df_admin["ФИО"].str.contains(search_query, case=False, na=False)]
            else:
                df_filtered = df_admin
                
            st.dataframe(df_filtered, use_container_width=True)
            
            st.write("### 📥 Выгрузить объективку сотрудника (МАЪЛУМОТНОМА):")
            employee_list = df_admin["ФИО"].tolist()
            selected_emp = st.selectbox("Выберите сотрудника для выгрузки:", employee_list)
            
            if selected_emp:
                selected_row = df_admin[df_admin["ФИО"] == selected_emp].iloc[0].to_dict()
                docx_buffer = generate_docx_anketa(selected_row)
                
                st.download_button(
                    label=f"📄 Танлаш: {selected_emp}.docx",
                    data=docx_buffer,
                    file_name=f"Malumotnoma_{selected_emp.replace(' ', '_')}.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document"
                )
            
            st.write("---")
            with open(FILE_NAME, "rb") as file:
                st.download_button(
                    label=T["admin_download"],
                    data=file,
                    file_name="Сотрудники_База_Данных.xlsx",
                    mime="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
                )
        else:
            st.info(T["admin_empty"])
            
    elif input_password != "":
        st.error(T["admin_err"])
